import os

import fitz  # PyMuPDF
import pytesseract
from docx import Document as DocxDocument
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation


def _extract_pdf(path: str) -> str:
    text = ""
    doc = fitz.open(path)

    for page in doc:
        page_text = page.get_text().strip()

        if page_text:
            text += page_text + "\n"
        else:
            # No extractable text layer - this is very likely a scanned
            # page. Render it to an image and OCR it instead of silently
            # returning nothing, which is what the original implementation
            # did for every scanned document.
            pix = page.get_pixmap(dpi=200)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            ocr_text = pytesseract.image_to_string(img)
            text += ocr_text + "\n"

    doc.close()
    return text


def _extract_docx(path: str) -> str:
    doc = DocxDocument(path)
    return "\n".join(p.text for p in doc.paragraphs)


def _extract_xlsx(path: str) -> str:
    workbook = load_workbook(path, data_only=True)
    text = ""
    for sheet in workbook:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value:
                    text += str(cell.value) + " "
    return text


def _extract_pptx(path: str) -> str:
    presentation = Presentation(path)
    text = ""
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text


def _extract_image(path: str) -> str:
    # Scanned single-page documents saved as an image rather than a PDF.
    return pytesseract.image_to_string(Image.open(path))


def extract_text(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()

    try:
        if ext == ".pdf":
            return _extract_pdf(file_path)
        if ext in (".docx", ".doc"):
            return _extract_docx(file_path)
        if ext in (".xlsx", ".xls"):
            return _extract_xlsx(file_path)
        if ext == ".pptx":
            return _extract_pptx(file_path)
        if ext in (".png", ".jpg", ".jpeg"):
            return _extract_image(file_path)
        if ext == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
    except Exception as exc:
        return f"[Could not extract text from this file: {exc}]"

    return ""


def create_chunks(text: str, chunk_size: int = 400) -> list:
    words = text.split()
    if not words:
        return []
    return [" ".join(words[i:i + chunk_size]) for i in range(0, len(words), chunk_size)]
