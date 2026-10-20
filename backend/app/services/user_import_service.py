"""
Reads a spreadsheet (CSV or Excel) of people and checks every row, so an
administrator can create many accounts at once.

Columns (names are matched loosely, in any order):
    ec_number   required
    full_name   required
    username    optional - made from the name when blank (e.g. "tmoyo")
    department  optional - must already exist, unless "create departments" is ticked
    role        optional - User (default), Manager or Admin
"""
import csv
import io
import re
import secrets
from typing import List, Optional

from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.database import Department, User
from app.services.department_service import find_department

MAX_ROWS = 300
MAX_FILE_BYTES = 2 * 1024 * 1024
ROLES = {"user": "User", "employee": "User", "manager": "Manager", "admin": "Admin"}

_HEADER_ALIASES = {
    "ec_number": {"ec", "ecnumber", "ecno", "ecnum", "ecnumbers", "employeenumber", "staffnumber"},
    "full_name": {"fullname", "name", "names", "employeename", "staffname", "fullnames"},
    "username": {"username", "user", "login", "userid"},
    "department": {"department", "dept", "departments", "section"},
    "role": {"role", "type", "accesslevel"},
}
_TEMP_CHARS = "ABCDEFGHJKMNPQRSTUVWXYZabcdefghijkmnpqrstuvwxyz23456789"


def temp_password() -> str:
    return "".join(secrets.choice(_TEMP_CHARS) for _ in range(10))


def _norm_header(h) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    text = str(v).strip()
    if re.fullmatch(r"\d+\.0", text):   # "1015.0" from a spreadsheet export
        return text[:-2]
    return text


def _read_rows(filename: str, data: bytes) -> List[List[str]]:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            rows = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
            wb.close()
            return rows
        except Exception:
            raise HTTPException(status_code=400, detail="That Excel file couldn't be read. Save it as .xlsx or .csv and try again.")
    if name.endswith(".csv") or name.endswith(".txt"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    raise HTTPException(status_code=400, detail="Upload a .csv or .xlsx file.")


def _make_username(full_name: str, taken: set) -> str:
    parts = re.sub(r"[^A-Za-z\s'-]", "", full_name).replace("'", "").replace("-", "").split()
    if not parts:
        base = "user"
    elif len(parts) == 1:
        base = parts[0].lower()
    else:
        base = (parts[0][0] + parts[-1]).lower()
    candidate, n = base, 1
    while candidate in taken:
        n += 1
        candidate = f"{base}{n}"
    return candidate


def check_file(db: Session, filename: str, data: bytes, create_departments: bool) -> dict:
    """Validate every row without creating anything."""
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(status_code=400, detail="That file is too large (2 MB at most).")
    raw = _read_rows(filename, data)
    raw = [r for r in raw if any(c for c in r)]
    if not raw:
        raise HTTPException(status_code=400, detail="The file is empty.")

    # map columns
    col = {}
    for i, h in enumerate(raw[0]):
        key = _norm_header(h)
        for field, aliases in _HEADER_ALIASES.items():
            if key in aliases and field not in col:
                col[field] = i
    missing = [f for f in ("ec_number", "full_name") if f not in col]
    if missing:
        raise HTTPException(
            status_code=400,
            detail="The first row must be column headings including: ec_number and full_name. "
                   "Use the template if you're unsure.",
        )

    body = raw[1:]
    if not body:
        raise HTTPException(status_code=400, detail="There are no people below the headings.")
    if len(body) > MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"Too many rows ({len(body)}). Import at most {MAX_ROWS} at a time.")

    taken_usernames = {u[0].lower() for u in db.query(User.username).all()}
    taken_ecs = {e[0].lower() for e in db.query(User.ec_number).all()}
    known_depts = {d.name.lower(): d.name for d in db.query(Department).all()}

    seen_ecs, seen_names = set(), set()
    out, new_departments = [], []

    def get(row, field):
        i = col.get(field)
        return row[i].strip() if i is not None and i < len(row) else ""

    for n, row in enumerate(body, start=2):  # row 1 is the headings
        ec = get(row, "ec_number")
        name = get(row, "full_name")
        username = get(row, "username")
        dept_in = get(row, "department")
        role_in = get(row, "role")
        problems = []

        if not ec:
            problems.append("EC number is missing")
        elif ec.lower() in taken_ecs:
            problems.append("EC number already has an account")
        elif ec.lower() in seen_ecs:
            problems.append("EC number appears twice in the file")
        if not name:
            problems.append("name is missing")

        role = ROLES.get(role_in.lower()) if role_in else "User"
        if role is None:
            problems.append(f"role '{role_in}' isn't User, Manager or Admin")

        dept = ""
        if dept_in:
            found = known_depts.get(dept_in.lower())
            if found:
                dept = found
            elif create_departments:
                dept = dept_in
                if dept_in.lower() not in {d.lower() for d in new_departments}:
                    new_departments.append(dept_in)
            else:
                problems.append(f"department '{dept_in}' doesn't exist")

        if username:
            u = username.lower()
            if u in taken_usernames or u in seen_names:
                problems.append("username is already taken")
        elif name:
            username = _make_username(name, taken_usernames | seen_names)

        if not problems:
            seen_ecs.add(ec.lower())
            seen_names.add(username.lower())

        out.append({
            "row": n, "ec_number": ec, "full_name": name, "username": username,
            "department": dept, "role": role or "", "status": "error" if problems else "ok",
            "message": "; ".join(problems),
        })

    return {
        "total": len(out),
        "valid": sum(1 for r in out if r["status"] == "ok"),
        "rows": out,
        "new_departments": new_departments,
    }
